import csv
import io
from urllib.parse import quote

from fastapi.responses import StreamingResponse
from openpyxl import Workbook


def _attachment_headers(filename: str) -> dict[str, str]:
    quoted = quote(filename)
    return {"Content-Disposition": f"attachment; filename=\"{quoted}\"; filename*=UTF-8''{quoted}"}


def rows_to_csv(headers: list[str], rows: list[list], filename: str) -> StreamingResponse:
    buffer = io.StringIO()
    # 加 BOM 让 Excel 直接识别 UTF-8 中文。
    buffer.write("﻿")
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.getvalue().encode("utf-8")]),
        media_type="text/csv; charset=utf-8",
        headers=_attachment_headers(filename),
    )


def tables_to_xlsx(sheets: dict[str, tuple[list[str], list[list]]], filename: str) -> StreamingResponse:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, (headers, rows) in sheets.items():
        sheet = workbook.create_sheet(title=sheet_name[:31] or "Sheet1")
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

    if not workbook.sheetnames:
        workbook.create_sheet(title="Sheet1")

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_headers(filename),
    )
